import hashlib
import io
import ipaddress
import socket
import zipfile
from collections.abc import Iterator
from urllib.parse import urlsplit
from xml.etree import ElementTree

import httpx
from docx import Document as DocxDocument
from docx.table import Table
from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from pypdf import PdfReader

from patch_ai.adapters import ocr
from patch_ai.api.budget import remaining_seconds
from patch_ai.config import Settings
from patch_ai.schemas.contracts import SourceFile


class SourceRejected(ValueError):
    """Safe, stable download/parser failure; never include URLs or source text."""


def validate_source_url(url: str, settings: Settings) -> None:
    parsed = urlsplit(url)
    if (
        parsed.scheme != "https"
        or parsed.hostname not in settings.allowed_source_hosts
        or parsed.username
        or parsed.password
        or parsed.port not in (None, 443)
        or parsed.fragment
    ):
        raise SourceRejected("SOURCE_URL_REJECTED")
    # A deployment egress firewall must also deny private ranges to prevent DNS rebinding.
    addresses = socket.getaddrinfo(parsed.hostname, 443, type=socket.SOCK_STREAM)
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise SourceRejected("SOURCE_URL_REJECTED")


def download_source(source: SourceFile, settings: Settings) -> bytes:
    validate_source_url(str(source.url), settings)
    if source.content_type not in settings.supported_mime_types:
        raise SourceRejected("SOURCE_TYPE_REJECTED")
    body = bytearray()
    with (
        httpx.Client(
            timeout=remaining_seconds(settings.source_download_timeout_seconds),
            follow_redirects=False,
            trust_env=False,
        ) as client,
        client.stream("GET", str(source.url)) as response,
    ):
        if response.status_code != 200:
            raise SourceRejected("SOURCE_DOWNLOAD_FAILED")
        if response.headers.get("content-type", "").split(";")[0] != source.content_type:
            raise SourceRejected("SOURCE_TYPE_MISMATCH")
        for chunk in response.iter_bytes(65536):
            remaining_seconds(settings.source_download_timeout_seconds)
            body.extend(chunk)
            if len(body) > settings.source_download_max_bytes:
                raise SourceRejected("SOURCE_TOO_LARGE")
    if not body or hashlib.sha256(body).hexdigest() != source.sha256.lower():
        raise SourceRejected("SOURCE_CHECKSUM_MISMATCH")
    return bytes(body)


class VerifiedSourceLoader(BaseLoader):
    """Local parser exception documented in ai/adapters/README.md."""

    def __init__(self, data: bytes, content_type: str, settings: Settings | None = None) -> None:
        self.data = data
        self.content_type = content_type
        self.settings = settings if settings is not None else Settings.model_construct()

    def lazy_load(self) -> Iterator[Document]:
        if (
            self.content_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            yield from self.docx_blocks()
            return
        if self.content_type in {"text/plain", "text/markdown"}:
            text = self.data.decode("utf-8-sig", errors="strict").strip()
            if not text or "\x00" in text or len(text) > 2_000_000:
                raise SourceRejected("SOURCE_TEXT_INVALID")
            for index, offset in enumerate(range(0, len(text), 20000), start=1):
                yield Document(
                    page_content=text[offset : offset + 20000],
                    metadata={
                        "page": index,
                        "section": f"Text block {index}",
                        "extractionQuality": 1.0,
                    },
                )
            return
        if self.content_type != "application/pdf" or not self.data.startswith(b"%PDF-"):
            raise SourceRejected("SOURCE_TYPE_REJECTED")
        reader = PdfReader(io.BytesIO(self.data), strict=True)
        if reader.is_encrypted or not 1 <= len(reader.pages) <= 500:
            raise SourceRejected("SOURCE_PDF_REJECTED")
        total = 0
        for index, page in enumerate(reader.pages, start=1):
            remaining_seconds(15)
            text = (page.extract_text() or "").strip()
            quality = 1.0
            if len(text) < 30:
                with ocr.render_page(self.data, index - 1, self.settings) as rendered:
                    text = ocr.image_text(rendered, self.settings)
                quality = 0.6  # OCR is explicitly reviewable, never equivalent to native text.
            elif len(page.images):
                # Native headings must not suppress text inside raster diagrams.
                if len(page.images) > 100:
                    raise SourceRejected("SOURCE_IMAGE_LIMIT_EXCEEDED")
                texts = [text]
                for embedded in page.images:
                    if embedded.image is None:
                        raise SourceRejected("SOURCE_IMAGE_UNREADABLE")
                    raster_text = ocr.image_text(embedded.image, self.settings)
                    if raster_text and raster_text not in text:
                        texts.append(raster_text)
                text = "\n\n".join(texts)
                quality = 0.6
            total += len(text)
            if not text or len(text) > 200_000 or total > 2_000_000:
                raise SourceRejected("SOURCE_PAGE_UNREADABLE")
            yield Document(
                page_content=text,
                metadata={"page": index, "section": f"Page {index}", "extractionQuality": quality},
            )

    def docx_blocks(self) -> Iterator[Document]:
        # Never extract archive members to disk or follow external relationships.
        with zipfile.ZipFile(io.BytesIO(self.data)) as package:
            entries = package.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 100_000_000:
                raise SourceRejected("SOURCE_ARCHIVE_TOO_LARGE")
            for entry in entries:
                if entry.flag_bits & 1 or entry.file_size > max(entry.compress_size, 1) * 100:
                    raise SourceRejected("SOURCE_ARCHIVE_REJECTED")
                if entry.filename.startswith(("word/media/", "word/embeddings/")):
                    raise SourceRejected("SOURCE_DOCX_VISUAL_CONTENT_REQUIRES_PDF")
                if entry.filename.endswith((".xml", ".rels")):
                    xml = package.read(entry).upper()
                    if any(
                        marker in xml
                        for marker in (
                            b"<!DOCTYPE",
                            b"<!ENTITY",
                            b"<W:INS ",
                            b"<W:DEL ",
                            b"<W:TXBXCONTENT",
                            b'TARGETMODE="EXTERNAL"',
                        )
                    ):
                        raise SourceRejected("SOURCE_DOCX_UNSUPPORTED_CONTENT")
                    root = ElementTree.fromstring(package.read(entry))
                    for element in root.iter():
                        tag = element.tag.rsplit("}", 1)[-1]
                        if (
                            tag
                            in {
                                "ins",
                                "del",
                                "txbxContent",
                                "drawing",
                                "pict",
                                "object",
                                "altChunk",
                                "footnoteReference",
                                "endnoteReference",
                                "moveFrom",
                                "moveTo",
                                "oMath",
                                "oMathPara",
                                "fldChar",
                            }
                            or element.attrib.get("TargetMode", "").upper() == "EXTERNAL"
                        ):
                            raise SourceRejected("SOURCE_DOCX_UNSUPPORTED_CONTENT")
        document = DocxDocument(io.BytesIO(self.data))
        texts = []
        for block in document.iter_inner_content():
            if isinstance(block, Table):
                if any(cell.tables for row in block.rows for cell in row.cells):
                    raise SourceRejected("SOURCE_DOCX_NESTED_TABLE_REQUIRES_PDF")
                texts.append(
                    "\n".join(" | ".join(cell.text for cell in row.cells) for row in block.rows)
                )
            else:
                texts.append(block.text)
        for section in document.sections:
            for part in (section.header, section.footer):
                texts.extend(p.text for p in part.paragraphs)
                if part.tables:
                    raise SourceRejected("SOURCE_DOCX_HEADER_TABLE_REQUIRES_PDF")
        text = "\n".join(texts).strip()
        if not text or len(text) > 2_000_000 or "\x00" in text:
            raise SourceRejected("SOURCE_TEXT_INVALID")
        for index, offset in enumerate(range(0, len(text), 20000), 1):
            yield Document(
                page_content=text[offset : offset + 20000],
                metadata={
                    "page": index,
                    "section": f"DOCX text block {index}",
                    "extractionQuality": 0.9,
                },
            )
