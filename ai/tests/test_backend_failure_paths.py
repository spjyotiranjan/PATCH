import asyncio
import io
import json
import threading
import time
import zipfile
from typing import Any
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from test_backend_workflows import FixtureProviders, fixtures

from patch_ai.adapters.providers import Model
from patch_ai.adapters.source_loader import SourceRejected, VerifiedSourceLoader
from patch_ai.api.budget import remaining_seconds, workflow_deadline
from patch_ai.api.execution import WorkflowCapacityError, WorkflowExecutor
from patch_ai.schemas.contracts import LogDraftRequest, LogScope, RevalidationRequest
from patch_ai.services.drafting import (
    FormattedLog,
    ProcedureAssessment,
    StepEvidenceVerification,
    draft_log,
    revalidate,
    review_analysis,
)

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_expired_workflow_cannot_start_more_provider_work() -> None:
    marker = workflow_deadline.set(time.monotonic() - 1)
    try:
        with pytest.raises(TimeoutError, match="WORKFLOW_DEADLINE_EXCEEDED"):
            remaining_seconds(30)
    finally:
        workflow_deadline.reset(marker)
    assert remaining_seconds(30) == 30


def docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_paragraph("Synthetic test card")
    row = document.add_table(rows=1, cols=2).rows[0]
    row.cells[0].text = "Label"
    row.cells[1].text = "Amber"
    document.add_paragraph("End of card")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_docx_preserves_body_table_order_and_block_anchors() -> None:
    pages = VerifiedSourceLoader(docx_bytes(), DOCX).load()
    assert pages[0].page_content == "Synthetic test card\nLabel | Amber\nEnd of card"
    assert pages[0].metadata == {
        "page": 1,
        "section": "DOCX text block 1",
        "extractionQuality": 0.9,
    }


@pytest.mark.parametrize(
    "filename,content",
    [
        ("word/media/image.png", b"fake image"),
        ("word/embeddings/object.bin", b"fake embedded file"),
        ("custom.xml", b'<!DOCTYPE root [<!ENTITY hidden "value">]><root/>'),
        (
            "custom.xml",
            b'<root xmlns:x="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><x:ins/></root>',
        ),
        (
            "custom.rels",
            b"<Relationships><Relationship TargetMode='External' Target='https://example.invalid'/></Relationships>",
        ),
        ("word/bomb.txt", b"a" * 100000),
    ],
    ids=["image", "embedded-object", "doctype", "tracked-change", "external-link", "zip-bomb"],
)
def test_docx_rejects_omitted_visuals_external_xml_and_archive_bombs(
    filename: str, content: bytes
) -> None:
    buffer = io.BytesIO(docx_bytes())
    with zipfile.ZipFile(buffer, "a", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(filename, content)
    with pytest.raises(SourceRejected):
        VerifiedSourceLoader(buffer.getvalue(), DOCX).load()


class ReviewProviders(FixtureProviders):
    critical = False
    assessment_critical = False

    def model(
        self,
        schema: type[Model],
        system: str,
        data: str,
        *,
        routing: bool = False,
        complex_reasoning: bool = False,
        images: tuple[bytes, ...] = (),
    ) -> Model:
        assert "untrusted" in system
        payload: dict[str, Any]
        if schema is StepEvidenceVerification:
            sources = json.loads(data)["sources"]
            payload = {
                "supported": True,
                "highCriticality": self.critical,
                "conflict": self.conflict,
                "missingMandatorySafetyEvidence": False,
                "chunkIds": ["unknown" if self.unsupported else sources[0]["chunk_id"]],
            }
        elif schema is ProcedureAssessment:
            payload = {
                "coverageComplete": True,
                "conflict": self.conflict,
                "applicable": True,
                "highCriticality": self.assessment_critical,
                "missingMandatorySafetyEvidence": False,
                "requiredTopics": ["label"],
                "missingTopics": [],
            }
        elif schema is FormattedLog:
            payload = {"text": data}
        else:
            return super().model(schema, system, data, routing=routing)
        return schema.model_validate(payload)


@pytest.mark.parametrize(
    "critical,assessment_critical", [(False, False), (True, False), (False, True)]
)
def test_revalidation_binds_each_unchanged_step_and_rejects_unknown_chunks(
    critical: bool,
    assessment_critical: bool,
) -> None:
    settings, fixture, question = fixtures()
    providers = ReviewProviders(settings)
    providers.critical = critical
    providers.assessment_critical = assessment_critical
    providers.records = fixture.records
    request = RevalidationRequest.model_validate(
        {
            "requestId": str(uuid4()),
            "contractVersion": "v1",
            "tenantId": "tenant-1",
            "projectId": "project-1",
            "projectDescription": "Synthetic test card review",
            "generationRequestId": "generation-1",
            "inputFingerprint": "a" * 64,
            "timezone": "Asia/Kolkata",
            "retrievalScopeManifest": question.retrieval_scope_manifest.model_dump(by_alias=True),
            "activeSources": [
                {
                    "documentVersionId": "version-1",
                    "documentTitle": "Test card",
                    "revision": "1",
                    "inclusionPath": "PROJECT_DIRECT",
                }
            ],
            "steps": [
                {
                    "stepId": "stable-step",
                    "position": 1,
                    "title": "Test label",
                    "instructions": "The test display label is amber.",
                    "required": True,
                    "citationIds": [],
                    "evidenceState": "UNVERIFIED",
                }
            ],
        }
    )
    before = request.model_dump()
    result = revalidate(request, settings, providers)
    assert result.status == "validated", result
    assert result.review_analysis.review_need == (
        "HIGH" if critical or assessment_critical else "LOW"
    )
    assert result.review_analysis.hardware_criticality == (
        "HIGH" if critical or assessment_critical else "NORMAL"
    )
    assert result.supported_step_ids == ["stable-step"]
    assert result.step_citations[0].citation_ids == [result.citations[0].id]
    assert request.model_dump() == before
    providers.unsupported = True
    result = revalidate(request, settings, providers)
    assert result.status == "needs_review" and not result.supported_step_ids
    assert not result.step_citations
    providers.conflict = True
    result = revalidate(request, settings, providers)
    assert result.review_analysis.review_need == "SEVERE"
    assert "CURRENT_SOURCE_CONFLICT" in result.review_analysis.blocking_findings


@pytest.mark.parametrize(
    "factor,expected",
    [
        (None, "LOW"),
        ("applicable", "MODERATE"),
        ("coverage_complete", "HIGH"),
        ("critical", "HIGH"),
        ("conflict", "SEVERE"),
        ("current", "SEVERE"),
        ("safety_gap", "SEVERE"),
    ],
)
def test_review_rubric_never_averages_away_a_severe_finding(
    factor: str | None, expected: str
) -> None:
    values = dict(
        coverage_complete=True,
        conflict=False,
        current=True,
        applicable=True,
        critical=False,
        safety_gap=False,
    )
    if factor:
        values[factor] = not values[factor]
    assert review_analysis(**values).review_need == expected


def test_log_helper_remains_draft_and_rejects_unverified_observations() -> None:
    settings, _, _ = fixtures()
    providers = ReviewProviders(settings)
    request = LogDraftRequest(
        request_id=uuid4(),
        contract_version="v1",
        project_id="project-1",
        scope_type=LogScope.PROJECT,
        source_text="Observed the synthetic amber label.",
    )
    result = draft_log(request, providers)
    assert result.status == "drafted" and result.draft_text == request.source_text
    assert "Editable draft; user submission required." in result.warnings
    providers.unsupported = True
    assert draft_log(request, providers).status == "unavailable"


@pytest.mark.asyncio
async def test_timed_out_work_retains_capacity_until_thread_finishes() -> None:
    executor = WorkflowExecutor(capacity=1)
    started, finish = threading.Event(), threading.Event()

    def blocked() -> str:
        started.set()
        finish.wait(2)
        return "done"

    try:
        with pytest.raises(WorkflowCapacityError):
            await executor.run(blocked, 0)
        assert await asyncio.to_thread(started.wait, 1)
        with pytest.raises(WorkflowCapacityError):
            await executor.run(lambda: "must not start", 1)
    finally:
        finish.set()
    for _ in range(100):
        if not executor.semaphore.locked():
            break
        await asyncio.sleep(0.01)
    assert await executor.run(lambda: "available", 1) == "available"
